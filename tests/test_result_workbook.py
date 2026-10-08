"""Independent anonymous XLSX/CSV fixtures, not model-generated answers."""
import csv
import builtins
import ctypes
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("check_result_workbook", ROOT / "scripts/check_result_workbook.py")
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        # Fixtures and hooks use the same canonical spelling as the checker;
        # Windows runners may expose TEMP through a RUNNER~1 alias.
        self.root = Path(self.temp.name).resolve(strict=True)
        self.template = self.root / "template.xlsx"
        self.result = self.root / "result.xlsx"
        self.body = self.root / "body.csv"
        self.contract = self.root / "contract.json"
        self.headers = ["time", "location_a", "location_b", "surface"]
        self.rows = [[t, 1 - 0.2 * t, 0.9 - 0.18 * t, 0.8 - 0.17 * t] for t in range(5)]
        self.workbook(self.template, [])
        self.workbook(self.result, self.rows)
        self.csv([["time", "a", "b", "surface"], [0, 1, 0.9, 0.8], [2, 0.6, 0.54, 0.46], [4, 0.2, 0.18, 0.12]])
        self.cfg = {"schema_version": 1,
                    "result": {"path": "result.xlsx", "sha256": self.sha(self.result), "value_precision": "unrounded"},
                    "template": {"path": "template.xlsx", "sha256": self.sha(self.template)},
                    "sheets": [{"name": "Series", "time_column": 1, "start": 0, "step": 1,
                                "minimum_end": 4, "allowed_final_event": False, "data_columns": [2, 3, 4],
                                "first_row_values": {"2": 1, "3": 0.9, "4": 0.8}, "max_value_exclusive": 0.25,
                                "body_samples": {"table": "body.csv", "table_sha256": self.sha(self.body), "time_key": "time",
                                                 "column_mappings": [{"body_column": "a", "workbook_column": 2, "scale": 1, "offset": 0},
                                                                     {"body_column": "b", "workbook_column": 3, "scale": 1, "offset": 0},
                                                                     {"body_column": "surface", "workbook_column": 4, "scale": 1, "offset": 0}],
                                                 "tolerance": 1e-8}}],
                    "output_paths": ["result.xlsx", "body.csv"]}
        self.save()

    @staticmethod
    def sha(path):
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def workbook(self, path, rows, headers=None, name="Series"):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = name
        sheet.append(headers or self.headers)
        for row in rows:
            sheet.append(row)
        book.save(path)
        book.close()

    def csv(self, rows):
        with self.body.open("w", encoding="utf-8", newline="") as stream:
            csv.writer(stream).writerows(rows)

    def save(self):
        self.contract.write_text(json.dumps(self.cfg), encoding="utf-8")

    def update_result(self, rows):
        self.workbook(self.result, rows)
        self.cfg["result"]["sha256"] = self.sha(self.result)
        self.save()

    def run_check(self):
        self.save()
        return m.check(self.contract, self.root)

    def assert_failed(self, substring):
        report = self.run_check()
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any(substring in error for error in report["errors"]), report["errors"])
        return report

    def test_complete_one_second_series_matches_template_and_body_samples_without_writes(self):
        before = {p.name: self.sha(p) for p in self.root.iterdir()}
        report = m.check(self.contract, self.root)
        self.assertEqual(report["status"], "mechanical pass", report["errors"])
        self.assertFalse(report["scientific_verified"])
        self.assertTrue(all(item["unchanged"] for item in report["input_hashes"]))
        self.assertEqual(before, {p.name: self.sha(p) for p in self.root.iterdir()})
        self.assertTrue(any("complete-delivery" in text for text in report["not_covered"]))

    def test_first_three_hours_cannot_satisfy_longer_declared_process(self):
        self.cfg["sheets"][0]["minimum_end"] = 10801
        self.cfg["sheets"][0].pop("body_samples")
        self.cfg["sheets"][0].pop("max_value_exclusive")
        rows = [[t, 1, 0.9, 0.8] for t in range(10801)]
        self.update_result(rows)
        report = self.assert_failed("minimum_end")
        self.assertTrue(any(c["passed"] and "every declared step" in c["check"] for c in report["checks"]))

    def test_600_second_sparse_rows_cannot_pass_one_second_contract(self):
        self.cfg["sheets"][0]["minimum_end"] = 1800
        self.cfg["sheets"][0].pop("body_samples")
        self.cfg["sheets"][0].pop("max_value_exclusive")
        self.update_result([[t, 1, 0.9, 0.8] for t in (0, 600, 1200, 1800)])
        self.assert_failed("every declared step")

    def test_missing_duplicate_and_reversed_time_rows_are_rejected(self):
        for rows, expected in ((self.rows[:2] + self.rows[3:], "every declared step"),
                               (self.rows[:2] + [self.rows[1]] + self.rows[2:], "strictly increasing"),
                               ([self.rows[0], self.rows[2], self.rows[1], *self.rows[3:]], "strictly increasing")):
            self.update_result(rows)
            self.assert_failed(expected)

    def test_only_final_short_event_step_is_allowed(self):
        self.cfg["sheets"][0]["allowed_final_event"] = True
        self.cfg["sheets"][0]["minimum_end"] = 4.25
        self.cfg["sheets"][0].pop("body_samples")
        self.update_result(self.rows + [[4.25, 0.1, 0.09, 0.08]])
        self.assertEqual(self.run_check()["status"], "mechanical pass")
        self.update_result(self.rows[:2] + [[1.25, 0.7, 0.6, 0.5]] + self.rows[2:] + [[4.25, 0.1, 0.09, 0.08]])
        self.assert_failed("every declared step")
        self.update_result(self.rows + [[6.25, 0.1, 0.09, 0.08]])
        self.assert_failed("every declared step")

    def test_terminal_0150004_does_not_pass_strict_015_despite_display_rounding(self):
        self.cfg["sheets"][0]["max_value_exclusive"] = 0.15
        self.cfg["sheets"][0].pop("body_samples")
        rows = [*self.rows[:-1], [4, 0.150004, 0.1, 0.1]]
        self.update_result(rows)
        book = openpyxl.load_workbook(self.result)
        book["Series"]["B6"].number_format = "0.0000"
        book.save(self.result); book.close()
        self.cfg["result"]["sha256"] = self.sha(self.result)
        report = self.assert_failed("strictly below")
        actual = next(c for c in report["checks"] if "strictly below" in c["check"])
        self.assertEqual(actual["detail"]["maximum"], 0.150004)
        self.update_result([*self.rows[:-1], [4, 0.15, 0.1, 0.1]])
        self.assert_failed("strictly below")

    def test_rounded_or_missing_precision_cannot_prove_strict_threshold(self):
        for precision in ("rounded", "unspecified"):
            self.cfg["result"]["value_precision"] = precision
            self.assert_failed("unrounded input precision")
        self.cfg["result"].pop("value_precision")
        self.assert_failed("unrounded input precision")

    def test_wrong_spatial_mapping_and_kelvin_offset_are_detected(self):
        self.cfg["sheets"][0]["body_samples"]["column_mappings"][0]["workbook_column"] = 3
        self.cfg["sheets"][0]["body_samples"]["column_mappings"].pop(1)
        self.assert_failed("body sample")
        self.cfg["sheets"][0]["body_samples"]["column_mappings"] = [{"body_column": "a", "workbook_column": 2, "scale": 1, "offset": 273.15}]
        self.assert_failed("body sample")

    def test_explicit_scale_offset_can_match_units_without_inferring_them(self):
        self.csv([["time", "kelvin", "scaled"], [0, 274.15, 1000], [2, 273.75, 600], [4, 273.35, 200]])
        sample = self.cfg["sheets"][0]["body_samples"]
        sample["table_sha256"] = self.sha(self.body)
        sample["column_mappings"] = [{"body_column": "kelvin", "workbook_column": 2, "offset": -273.15},
                                     {"body_column": "scaled", "workbook_column": 3, "scale": 0.0009}]
        self.assertEqual(self.run_check()["status"], "mechanical pass")

    def test_wrong_surface_values_are_caught_by_first_row_and_body_samples(self):
        rows = [row[:] for row in self.rows]
        rows[0][3] = 1.8
        self.update_result(rows)
        self.assert_failed("first-row value column 4")
        rows = [row[:] for row in self.rows]
        rows[2][3] = 0.6
        self.update_result(rows)
        self.assert_failed("body sample row 2 column 4")

    def test_template_name_header_order_and_column_width_must_match(self):
        for headers, name in ((["time", "location_b", "location_a", "surface"], "Series"),
                              (self.headers + ["extra"], "Series"), (self.headers, "Renamed")):
            self.workbook(self.result, self.rows, headers, name)
            self.cfg["result"]["sha256"] = self.sha(self.result)
            self.assert_failed("match" if name == "Series" else "sheet")

    def expanded_fixture(self, surface=False):
        headers = ["time\\location", 0, 0.1, 0.2, "…", "surface" if surface else 2]
        expected = ["time\\location"] + [round(i / 10, 1) for i in range(21)] + (["surface"] if surface else [])
        self.workbook(self.template, [], headers)
        rows = [[t] + [1 - 0.2 * t] * (len(expected) - 1) for t in range(5)]
        actual_headers = [str(v) if isinstance(v, (int, float)) else v for v in expected]
        self.workbook(self.result, rows, actual_headers)
        source = self.root / "source.md"
        source.write_text("Anonymous task: locations from 0 to 2 at spacing 0.1; include surface when declared.\n", encoding="utf-8")
        self.cfg["result"]["sha256"] = self.sha(self.result)
        self.cfg["template"]["sha256"] = self.sha(self.template)
        sheet = self.cfg["sheets"][0]
        sheet.pop("body_samples")
        sheet["data_columns"] = list(range(2, len(expected) + 1))
        sheet["first_row_values"] = {"2": 1, str(len(expected)): 1}
        sheet["expected_header"] = expected
        sheet["header_basis"] = {"source": {"path": "source.md", "sha256": self.sha(source)},
                                 "location": "Task output-grid paragraph", "reason": "Expand the explicitly omitted intermediate locations"}
        return expected, rows, source

    def test_explicit_sha_bound_schematic_six_columns_can_expand_to_22_numeric_headers(self):
        expected, _, source = self.expanded_fixture()
        self.assertEqual(len(expected), 22)
        before = self.sha(source)
        report = self.run_check()
        self.assertEqual(report["status"], "mechanical pass", report["errors"])
        self.assertFalse(report["scientific_verified"])
        self.assertTrue(any("host declaration" in text and "expanded header" in text for text in report["not_covered"]))
        self.assertEqual(before, self.sha(source))
        self.assertTrue(any(item["path"] == "source.md" and item["unchanged"] for item in report["input_hashes"]))

    def test_explicit_schematic_surface_tail_can_expand_to_23_columns(self):
        expected, _, _ = self.expanded_fixture(surface=True)
        self.assertEqual(len(expected), 23)
        self.assertEqual(self.run_check()["status"], "mechanical pass")

    def test_schematic_expansion_requires_complete_basis_and_source_sha(self):
        self.expanded_fixture()
        sheet = self.cfg["sheets"][0]
        basis = json.loads(json.dumps(sheet["header_basis"]))
        sheet.pop("header_basis")
        self.assert_failed("header_basis")
        sheet["header_basis"] = basis
        sheet["header_basis"]["source"].pop("sha256")
        self.assert_failed("SHA-256 required")
        sheet["header_basis"] = json.loads(json.dumps(basis))
        sheet["header_basis"]["source"]["sha256"] = "0" * 64
        self.assert_failed("stale input SHA")
        sheet["header_basis"] = basis
        sheet["header_basis"]["reason"] = ""
        self.assert_failed("expansion reason")

    def test_no_automatic_schematic_expansion_without_explicit_header(self):
        self.expanded_fixture()
        sheet = self.cfg["sheets"][0]
        sheet.pop("expected_header")
        sheet.pop("header_basis")
        self.assert_failed("column count")

    def test_fixed_schematic_prefix_and_tail_cannot_be_rewritten_by_declared_expansion(self):
        expected, _, _ = self.expanded_fixture()
        sheet = self.cfg["sheets"][0]
        for index, wrong in ((0, "different time label"), (2, 0.11), (21, 3)):
            broken = expected[:]
            broken[index] = wrong
            sheet["expected_header"] = broken
            self.assert_failed("preserves original fixed")
        sheet["expected_header"] = expected[:5]
        self.assert_failed("preserves original fixed")

    def test_actual_wrong_coordinate_and_surface_tail_fail_expanded_header_check(self):
        expected, rows, _ = self.expanded_fixture(surface=True)
        for index, wrong in ((10, "0.95"), (22, "Surface")):
            changed = expected[:]
            changed[index] = wrong
            self.workbook(self.result, rows, changed)
            self.cfg["result"]["sha256"] = self.sha(self.result)
            self.assert_failed("header values match")

    def test_complete_original_template_cannot_be_overridden_even_if_expected_matches(self):
        source = self.root / "source.md"
        source.write_text("Host statement", encoding="utf-8")
        sheet = self.cfg["sheets"][0]
        sheet["expected_header"] = self.headers[:]
        sheet["header_basis"] = {"source": {"path": "source.md", "sha256": self.sha(source)},
                                 "location": "one paragraph", "reason": "attempt to replace a complete template"}
        self.assert_failed("without an ellipsis cannot be overridden")

    def test_numeric_header_strings_are_equivalent_but_text_whitespace_is_not_relaxed(self):
        self.assertTrue(m.header_equal("0.0", 0))
        self.assertTrue(m.header_equal("1.10", 1.1))
        self.assertFalse(m.header_equal("1.10001", 1.1))
        self.assertFalse(m.header_equal("surface ", "surface"))
        self.assertFalse(m.header_equal(True, 1))
        self.assertFalse(m.header_equal("0cm", 0))

    def test_ordered_fixed_blocks_between_multiple_ellipses_remain_anchored(self):
        original = ["time", 0, "...", 1, "…", 2]
        self.assertTrue(m.valid_expansion(original, ["time", 0, 0.5, 1, 1.5, 2]))
        self.assertFalse(m.valid_expansion(original, ["time", 0, 1.5, 2, 0.5, 1]))
        self.assertFalse(m.valid_expansion(original, ["time", 0, 1, 1.5, 2]))
        self.assertFalse(m.valid_expansion(original, ["time", 0, 0.5, 1, "…", 2]))

    def test_missing_csv_time_and_duplicate_body_time_fail(self):
        for rows in ([["time", "a", "b", "surface"], [1.5, 0.7, 0.63, 0.545]],
                     [["time", "a", "b", "surface"], [2, 0.6, 0.54, 0.46], [2, 0.6, 0.54, 0.46]]):
            self.csv(rows)
            self.cfg["sheets"][0]["body_samples"]["table_sha256"] = self.sha(self.body)
            self.assert_failed("body sample time")

    def test_casefold_output_path_collision_including_t_T_is_reported(self):
        self.cfg["output_paths"] = ["results/t.csv", "results/T.csv"]
        self.assert_failed("casefold collisions")

    def test_stale_or_missing_input_hashes_do_not_inspect_values(self):
        for field in ("result", "template"):
            original = self.cfg[field]["sha256"]
            self.cfg[field]["sha256"] = "0" * 64
            self.assert_failed("stale input SHA")
            self.cfg[field]["sha256"] = original
            self.cfg[field].pop("sha256")
            self.assert_failed("SHA-256 required")
            self.cfg[field]["sha256"] = original
        self.cfg["sheets"][0]["body_samples"]["table_sha256"] = "0" * 64
        self.assert_failed("stale input SHA")

    def test_hashes_are_checked_after_inspection_as_well_as_before(self):
        original = openpyxl.load_workbook
        triggered = False
        def load_and_tamper(path, **kwargs):
            nonlocal triggered
            book = original(path, **kwargs)
            if Path(path) == self.template:
                triggered = True
                self.body.write_text("time,a,b,surface\n0,1,0.9,0.8\n", encoding="utf-8")
            return book
        with patch.object(openpyxl, "load_workbook", side_effect=load_and_tamper):
            report = self.run_check()
        self.assertTrue(triggered, "Template-load tamper hook did not execute")
        self.assertEqual(report["status"], "failed")
        # CSV was changed before its declared binding; this must fail stale SHA.
        self.assertTrue(any("stale input SHA" in error for error in report["errors"]))

    def test_actual_after_read_input_change_fails_unchanged_check(self):
        original = m.sha256
        calls = 0
        triggered = False
        def alter_on_final_read(path):
            nonlocal calls, triggered
            if Path(path) == self.result:
                calls += 1
                if calls == 2:
                    triggered = True
                    self.result.write_bytes(self.result.read_bytes() + b"changed-after-inspection")
            return original(path)
        with patch.object(m, "sha256", side_effect=alter_on_final_read):
            report = m.check(self.contract, self.root)
        self.assertTrue(triggered, "After-read tamper hook did not execute")
        self.assertEqual(calls, 2, "Result must be hashed before and after inspection")
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any(not h["unchanged"] for h in report["input_hashes"] if h["path"] == "result.xlsx"))

    def test_traversal_absolute_paths_and_ancestor_links_are_refused(self):
        for value in ("../outside.xlsx", "C:/private.xlsx", "result.xlsx:stream", "/tmp/result.xlsx"):
            self.cfg["result"]["path"] = value
            self.assert_failed("Paths must stay relative")
        self.cfg["result"]["path"] = "result.xlsx"
        original = m.is_link
        result_triggered, ancestor_triggered = False, False
        def result_link(path):
            nonlocal result_triggered
            if Path(path) == self.result:
                result_triggered = True
                return True
            return original(path)
        def ancestor_link(path):
            nonlocal ancestor_triggered
            if Path(path) == self.root.parent:
                ancestor_triggered = True
                return True
            return original(path)
        with patch.object(m, "is_link", side_effect=result_link):
            self.assert_failed("junctions")
        self.assertTrue(result_triggered, "Result reparse-point hook did not execute")
        with patch.object(m, "is_link", side_effect=ancestor_link):
            self.assert_failed("root may not traverse")
        self.assertTrue(ancestor_triggered, "Root ancestor reparse-point hook did not execute")

    def test_absolute_contract_lexical_file_and_parent_links_are_refused_before_resolve(self):
        alias = self.root / "contract-alias.json"
        alias.write_bytes(self.contract.read_bytes())
        original = m.is_link
        with patch.object(m, "is_link", side_effect=lambda p: Path(p) == alias or original(p)):
            report = m.check(alias, self.root)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("Contract may not traverse" in e for e in report["errors"]))
        parent = self.root / "alias-parent"
        parent.mkdir()
        nested = parent / "contract.json"
        nested.write_bytes(self.contract.read_bytes())
        with patch.object(m, "is_link", side_effect=lambda p: Path(p) == parent or original(p)):
            report = m.check(nested, self.root)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("Contract may not traverse" in e for e in report["errors"]))

    def test_absolute_contract_outside_root_and_hardlinks_still_fail(self):
        with tempfile.TemporaryDirectory() as outside:
            path = Path(outside) / "contract.json"
            path.write_bytes(self.contract.read_bytes())
            report = m.check(path, self.root)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("inside project-root" in e for e in report["errors"]))
        alias = self.root / "contract-hardlink.json"
        os.link(self.contract, alias)
        report = m.check(alias, self.root)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("single-linked" in e for e in report["errors"]))

    def windows_short_path(self, path):
        if os.name != "nt":
            self.skipTest("Windows GetShortPathNameW is unavailable on this platform")
        function = ctypes.WinDLL("kernel32", use_last_error=True).GetShortPathNameW
        function.argtypes = [ctypes.c_wchar_p, ctypes.c_wchar_p, ctypes.c_uint32]
        function.restype = ctypes.c_uint32
        required = function(str(path), None, 0)
        if not required:
            self.skipTest("GetShortPathNameW unavailable for this filesystem: error " + str(ctypes.get_last_error()))
        buffer = ctypes.create_unicode_buffer(required + 1)
        actual = function(str(path), buffer, len(buffer))
        if not actual or actual >= len(buffer):
            self.skipTest("GetShortPathNameW could not supply a stable filesystem alias")
        short = Path(buffer.value)
        if str(short).casefold() == str(Path(path).resolve()).casefold():
            self.skipTest("This filesystem does not expose an 8.3 alias for the test path")
        self.assertEqual(short.resolve(strict=True), Path(path).resolve(strict=True))
        return short

    def test_windows_real_short_contract_with_long_root(self):
        # A long basename guarantees an alias where 8.3 generation is enabled,
        # even when the runner's temporary parent already uses a short name.
        long_contract = self.root / "anonymous workbook output verification contract.json"
        long_contract.write_bytes(self.contract.read_bytes())
        short_contract = self.windows_short_path(long_contract.resolve())
        report = m.check(short_contract, self.root.resolve())
        self.assertEqual(report["status"], "mechanical pass", report["errors"])

    def test_windows_real_long_contract_with_short_root(self):
        nested = self.root / "anonymous workbook verification project directory"
        nested.mkdir()
        for source in (self.template, self.result, self.body, self.contract):
            (nested / source.name).write_bytes(source.read_bytes())
        short_root = self.windows_short_path(nested.resolve())
        report = m.check((nested / "contract.json").resolve(), short_root)
        self.assertEqual(report["status"], "mechanical pass", report["errors"])

    def test_directories_hardlinks_and_invalid_xlsx_fail_closed(self):
        self.cfg["result"]["path"] = "."
        self.assert_failed("Ambiguous")
        self.cfg["result"]["path"] = "result.xlsx"
        other = self.root / "hardlink.xlsx"
        os.link(self.result, other)
        self.assert_failed("single-linked")
        other.unlink()
        self.result.write_bytes(b"not an XLSX")
        self.cfg["result"]["sha256"] = self.sha(self.result)
        self.assert_failed("XLSX could not be read")

    def test_formula_blank_and_nan_cells_do_not_pass_numeric_constant_contract(self):
        for value in ("=1+1", None, "nan"):
            rows = [row[:] for row in self.rows]
            rows[2][1] = value
            self.update_result(rows)
            self.assert_failed("finite numeric constants")

    def test_no_optional_claims_are_silently_added_and_unselected_columns_are_listed(self):
        self.cfg.pop("output_paths")
        sheet = self.cfg["sheets"][0]
        for key in ("first_row_values", "max_value_exclusive", "body_samples"):
            sheet.pop(key)
        sheet["data_columns"] = [2]
        report = self.run_check()
        self.assertEqual(report["status"], "mechanical pass")
        for phrase in ("case collisions", "initial physical", "threshold", "no body CSV", "unselected data columns"):
            self.assertTrue(any(phrase in text for text in report["not_covered"]))

    def test_unselected_domain_blanks_do_not_become_false_whole_domain_certification(self):
        rows = [row[:] for row in self.rows]
        for row in rows[2:]:
            row[3] = None
        self.update_result(rows)
        sheet = self.cfg["sheets"][0]
        sheet["data_columns"] = [2, 3]
        sheet["first_row_values"].pop("4")
        sheet["body_samples"]["column_mappings"].pop()
        report = self.run_check()
        self.assertEqual(report["status"], "mechanical pass", report["errors"])
        self.assertFalse(report["scientific_verified"])
        self.assertTrue(any("unselected data columns 4" in text for text in report["not_covered"]))
        self.assertTrue(any("all-domain terminal" in text for text in report["not_covered"]))
        threshold = next(c for c in report["checks"] if "strictly below" in c["check"])
        self.assertEqual(threshold["detail"]["checked_columns"], [2, 3])
        self.assertIn("selected columns only", threshold["detail"]["scope"])

    def test_missing_optional_dependency_is_reported_with_installation_target(self):
        original = builtins.__import__
        def without_openpyxl(name, *args, **kwargs):
            if name == "openpyxl":
                raise ImportError("controlled missing optional package")
            return original(name, *args, **kwargs)
        with patch.object(builtins, "__import__", side_effect=without_openpyxl):
            report = m.check(self.contract, self.root)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("Missing optional dependency openpyxl" in e and "requirements-results.txt" in e for e in report["errors"]))

    def test_cell_budget_failure_does_not_issue_partial_pass(self):
        with patch.object(m, "MAX_CELLS", 10):
            report = self.run_check()
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("no partial pass" in text for text in report["errors"]))

    def test_cli_writes_new_report_only_and_preserves_input_files(self):
        before = self.sha(self.result)
        cmd = [sys.executable, "-X", "utf8", str(ROOT / "scripts/check_result_workbook.py"), "--contract", str(self.contract),
               "--project-root", str(self.root), "--report", "report.json"]
        process = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertEqual(json.loads(process.stdout)["status"], "mechanical pass")
        self.assertFalse(json.loads((self.root / "report.json").read_text(encoding="utf-8"))["scientific_verified"])
        self.assertEqual(before, self.sha(self.result))
        original_report = (self.root / "report.json").read_bytes()
        again = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(again.returncode, 1)
        self.assertEqual((self.root / "report.json").read_bytes(), original_report)


if __name__ == "__main__":
    unittest.main()
