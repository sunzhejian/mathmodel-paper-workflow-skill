"""Read-only mechanical checks of a declared XLSX result/template/CSV contract.

No model execution, formula evaluation, unit inference or scientific approval is
performed. SHA-256 binds every data input before and after inspection. openpyxl is
an optional dependency supplied by requirements-results.txt.
"""
from __future__ import annotations

import argparse
import bisect
import csv
import hashlib
import json
import math
from pathlib import Path, PurePosixPath, PureWindowsPath
import re
import stat
import sys


MAX_CELLS = 20_000_000
ELLIPSIS_HEADERS = {"…", "...", "⋯", "……", "......"}


class ContractError(ValueError):
    pass


def sha256(path):
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def is_link(path):
    info = path.lstat()
    return stat.S_ISLNK(info.st_mode) or bool(getattr(info, "st_file_attributes", 0) &
                                            getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0))


def root_path(value):
    root = Path(value).absolute()
    for parent in (root, *root.parents):
        if parent.exists() and is_link(parent):
            raise ContractError("Project root may not traverse symlinks or junctions")
    if not root.is_dir():
        raise ContractError("project-root must be an existing directory")
    return root.resolve()


def normalized_path(value):
    if not isinstance(value, str) or not value or any(ord(c) < 32 for c in value):
        raise ContractError("A nonempty relative file path is required")
    value = value.replace("\\", "/")
    posix, windows = PurePosixPath(value), PureWindowsPath(value)
    if posix.is_absolute() or windows.drive or windows.root or ":" in value or ".." in posix.parts:
        raise ContractError("Paths must stay relative to project-root; drives, streams and traversal are forbidden")
    parts = tuple(part for part in posix.parts if part != ".")
    if not parts or any(part.endswith((".", " ")) for part in parts):
        raise ContractError("Ambiguous file path")
    return "/".join(parts)


def confined(root, value, *, existing=True):
    value = normalized_path(value)
    current = root
    for part in PurePosixPath(value).parts:
        current = current / part
        if current.exists() or current.is_symlink():
            if is_link(current):
                raise ContractError("Symlinks, junctions and reparse points are forbidden")
    if not current.resolve(strict=False).is_relative_to(root):
        raise ContractError("Path escapes project-root")
    if existing and not current.is_file():
        raise ContractError("Declared input is not an existing regular file")
    if current.exists() and (not current.is_file() or current.stat().st_nlink > 1):
        raise ContractError("Only ordinary single-linked files are accepted")
    return current


def object_fields(value, allowed, where):
    if not isinstance(value, dict) or set(value) - set(allowed):
        raise ContractError(where + ": object required; unknown fields are forbidden")


def number(value, where):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ContractError(where + ": finite numeric value required")
    return float(value)


def column(value, where):
    if type(value) is not int or value < 1:
        raise ContractError(where + ": positive one-based column required")
    return value


def ellipsis_header(value):
    return isinstance(value, str) and value.strip() in ELLIPSIS_HEADERS


def header_equal(actual, expected):
    """Numeric coordinate headers may be numbers or finite numeric strings."""
    def numeric(value):
        if isinstance(value, bool) or not isinstance(value, (str, int, float)):
            return None
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
        return number if math.isfinite(number) else None
    a, b = numeric(actual), numeric(expected)
    if a is not None and b is not None:
        return math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12)
    return type(actual) is type(expected) and actual == expected


def valid_expansion(original, expected):
    """Every ellipsis represents >=1 columns; fixed blocks stay ordered/anchored."""
    if len(expected) < len(original) or any(ellipsis_header(value) for value in expected):
        return False
    blocks = [[]]
    for value in original:
        if ellipsis_header(value):
            blocks.append([])
        else:
            blocks[-1].append(value)
    if len(blocks) < 2:
        return False
    def equal_block(position, block):
        return position >= 0 and position + len(block) <= len(expected) and all(
            header_equal(expected[position + index], value) for index, value in enumerate(block))
    if not equal_block(0, blocks[0]):
        return False
    suffix_start = len(expected) - len(blocks[-1])
    if not equal_block(suffix_start, blocks[-1]):
        return False
    cursor = len(blocks[0])
    for block in blocks[1:-1]:
        found = next((position for position in range(cursor + 1, suffix_start + 1)
                      if equal_block(position, block)), None)
        if found is None:
            return False
        cursor = found + len(block)
    return suffix_start >= cursor + 1


def read_contract(path):
    def unique(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ContractError("Duplicate contract JSON keys are forbidden")
            result[key] = value
        return result
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"), object_pairs_hook=unique,
                          parse_constant=lambda _: (_ for _ in ()).throw(ContractError("Nonfinite JSON numbers are forbidden")))
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise ContractError("Contract must be valid UTF-8 JSON") from exc


def check(contract_path, project_root):
    report = {"status": "failed", "scientific_verified": False, "checks": [], "errors": [], "input_hashes": [],
              "not_covered": ["Model equations, algorithms, physical correctness and solver execution are not verified.",
                              "Original precision is a host declaration; provenance and unroundedness are not independently proven.",
                              "Only declared XLSX rows, columns and CSV samples are checked; no complete-delivery approval is issued.",
                              "Moving-domain/radius-dependent blank legitimacy and all-domain terminal events are not verified. Selecting always-defined or surface columns does not prove whole-domain extrema or termination.",
                              "Formatting, figures, references, submission rules and undeclared artifacts are not checked."]}
    frozen, books = [], []
    def record(ok, name, detail=None):
        item = {"check": name, "passed": bool(ok)}
        if detail is not None:
            item["detail"] = detail
        report["checks"].append(item)
        if not ok:
            report["errors"].append(name)
    try:
        root = root_path(project_root)
        supplied = Path(contract_path)
        if supplied.is_absolute():
            lexical = supplied.absolute()
            # Canonicalize both Windows 8.3 aliases and long paths, but inspect
            # the original ancestry first: resolve() must not erase a link.
            for parent in (lexical, *lexical.parents):
                if (parent.exists() or parent.is_symlink()) and is_link(parent):
                    raise ContractError("Contract may not traverse symlinks, junctions or reparse points")
            canonical = lexical.resolve(strict=True)
            try:
                contract_key = canonical.relative_to(root).as_posix()
            except ValueError as exc:
                raise ContractError("Contract must also be inside project-root") from exc
        else:
            contract_key = normalized_path(str(supplied))
        contract_file = confined(root, contract_key)
        contract_hash = sha256(contract_file)
        frozen.append((contract_key, contract_file, contract_hash))
        spec = read_contract(contract_file)
        object_fields(spec, {"schema_version", "description", "result", "template", "sheets", "output_paths"}, "contract")
        if type(spec.get("schema_version")) is not int or spec["schema_version"] != 1:
            raise ContractError("Use contract schema_version 1")

        def bind(item, where, suffix):
            object_fields(item, {"path", "sha256", "value_precision"} if where == "result" else {"path", "sha256"}, where)
            key = normalized_path(item.get("path"))
            expected = item.get("sha256")
            if not isinstance(expected, str) or not re.fullmatch(r"[0-9a-f]{64}", expected):
                raise ContractError(where + ": explicit lowercase SHA-256 required")
            path = confined(root, key)
            if suffix is not None and path.suffix.lower() not in suffix:
                raise ContractError(where + ": unsupported file format")
            actual = sha256(path)
            frozen.append((key, path, actual))
            record(actual == expected, where + " SHA-256 matches", {"path": key, "sha256": actual})
            if actual != expected:
                raise ContractError(where + ": stale input SHA-256; no values inspected")
            return path

        result_path = bind(spec.get("result"), "result", {".xlsx"})
        template_path = bind(spec.get("template"), "template", {".xlsx"})
        if result_path == template_path:
            raise ContractError("Result and original template must be distinct files")
        precision = spec["result"].get("value_precision", "unspecified")
        if precision not in {"unrounded", "rounded", "unspecified"}:
            raise ContractError("result.value_precision must be unrounded, rounded or unspecified")
        report["value_precision_declaration"] = precision
        outputs = spec.get("output_paths")
        if outputs is None:
            report["not_covered"].append("Output filename case collisions were not declared for checking.")
        else:
            if not isinstance(outputs, list) or not all(isinstance(v, str) for v in outputs):
                raise ContractError("output_paths must be an explicit path array")
            keys = [normalized_path(v) for v in outputs]
            for key in keys:
                confined(root, key, existing=False)
            record(len({key.casefold() for key in keys}) == len(keys), "Output paths have no casefold collisions", {"paths": keys})
        try:
            import openpyxl
        except ImportError as exc:
            raise ContractError("Missing optional dependency openpyxl; install requirements-results.txt with this Python interpreter") from exc
        try:
            result = openpyxl.load_workbook(result_path, read_only=True, data_only=False, keep_links=False)
            books.append(result)
            template = openpyxl.load_workbook(template_path, read_only=True, data_only=False, keep_links=False)
            books.append(template)
        except Exception as exc:
            raise ContractError("Result/template XLSX could not be read; no partial pass is issued") from exc
        record(result.sheetnames == template.sheetnames, "Workbook sheet names/order match original template", {"result": result.sheetnames, "template": template.sheetnames})
        configs = spec.get("sheets")
        if not isinstance(configs, list) or not configs:
            raise ContractError("sheets must explicitly declare one or more sheet contracts")
        names = []
        for cfg in configs:
            object_fields(cfg, {"name", "time_column", "start", "step", "minimum_end", "allowed_final_event", "max_value_exclusive",
                                "data_columns", "first_row_values", "body_samples", "header_row", "time_tolerance", "tolerance",
                                "expected_header", "header_basis"}, "sheet")
            name = cfg.get("name")
            if not isinstance(name, str) or not name or name in names:
                raise ContractError("Sheet contract names must be distinct nonempty text")
            names.append(name)
            if name not in result.sheetnames or name not in template.sheetnames:
                record(False, name + ": declared sheet exists in result and template")
                continue
            sheet, original = result[name], template[name]
            time_col = column(cfg.get("time_column"), "time_column")
            header_row = column(cfg.get("header_row", 1), "header_row")
            start, step, minimum_end = (number(cfg.get(k), k) for k in ("start", "step", "minimum_end"))
            if step <= 0 or minimum_end < start:
                raise ContractError(name + ": positive step and minimum_end >= start are required")
            allowed_event = cfg.get("allowed_final_event", False)
            if type(allowed_event) is not bool:
                raise ContractError("allowed_final_event must be a boolean")
            time_tol = number(cfg.get("time_tolerance", max(1e-9, abs(step) * 1e-9)), "time_tolerance")
            tolerance = number(cfg.get("tolerance", 1e-8), "tolerance")
            if time_tol < 0 or tolerance < 0 or time_tol >= step / 2:
                raise ContractError("Tolerances must be nonnegative and time_tolerance less than half a step")
            template_width = original.max_column
            if header_row > original.max_row:
                raise ContractError(name + ": declared header row lies outside original template")
            template_header = next(original.iter_rows(min_row=header_row, max_row=header_row, max_col=template_width, values_only=True))
            expected_header = template_header
            expanded = "expected_header" in cfg or "header_basis" in cfg
            if expanded:
                if not any(ellipsis_header(value) for value in template_header):
                    raise ContractError(name + ": a complete original template without an ellipsis cannot be overridden")
                expected_header, basis = cfg.get("expected_header"), cfg.get("header_basis")
                if not isinstance(expected_header, list) or not expected_header or any(
                        value is not None and (isinstance(value, bool) or not isinstance(value, (str, int, float)) or
                        isinstance(value, (int, float)) and not math.isfinite(value)) for value in expected_header):
                    raise ContractError("expected_header must declare the complete scalar header without nonfinite numbers")
                object_fields(basis, {"source", "location", "reason"}, "header_basis")
                if not all(isinstance(basis.get(key), str) and basis[key].strip() for key in ("location", "reason")):
                    raise ContractError("header_basis requires explicit nonempty source location and expansion reason")
                source = bind(basis.get("source"), name + " header basis", None)
                record(valid_expansion(template_header, expected_header), name + ": explicit expanded header preserves original fixed prefix/suffix and ordered blocks")
                report["not_covered"].append(name + ": expanded header basis is a host declaration, not an automatic interpretation or verification of the original problem.")
                record(True, name + ": expanded header has an explicit SHA-bound source", {"source": source.relative_to(root).as_posix(),
                       "location": basis["location"], "reason": basis["reason"], "scientific_interpretation_verified": False})
            width = len(expected_header)
            if sheet.max_row * sheet.max_column > MAX_CELLS or original.max_row * template_width > MAX_CELLS:
                raise ContractError("Workbook rectangle exceeds the supported 20 million cells per sheet; no partial pass is issued")
            if time_col > width:
                raise ContractError(name + ": declared time column/header row lies outside original template")
            record(sheet.max_column == width, name + ": column count matches " + ("explicit expanded contract" if expanded else "original template"),
                   {"result": sheet.max_column, "expected": width, "original_template": template_width})
            got_header = next(sheet.iter_rows(min_row=header_row, max_row=header_row, max_col=width, values_only=True))
            record(all(header_equal(a, b) for a, b in zip(got_header, expected_header)), name + ": header values match " +
                   ("explicit expanded contract" if expanded else "original template"))
            columns = cfg.get("data_columns", [c for c in range(1, width + 1) if c != time_col])
            if not isinstance(columns, list) or not columns or len(set(columns)) != len(columns):
                raise ContractError("data_columns must contain distinct one-based columns")
            columns = [column(c, "data_columns") for c in columns]
            if time_col in columns or max(columns) > width:
                raise ContractError("data_columns must be non-time columns within the original template")
            ignored = [c for c in range(1, width + 1) if c not in columns and c != time_col]
            if ignored:
                report["not_covered"].append(name + ": unselected data columns " + ",".join(map(str, ignored)) + " are not exhaustively inspected.")
            rows = list(sheet.iter_rows(min_row=header_row + 1, max_col=width, values_only=True))
            while rows and all(value is None for value in rows[-1]):
                rows.pop()
            record(bool(rows), name + ": result contains data rows")
            if not rows:
                continue
            times, numeric_ok = [], True
            for row in rows:
                try:
                    times.append(number(row[time_col - 1], "time cell"))
                    for c in columns:
                        number(row[c - 1], "data cell")
                except ContractError:
                    numeric_ok = False
                    break
            record(numeric_ok, name + ": declared time/data cells are finite numeric constants (no formulas or blanks)")
            if not numeric_ok:
                continue
            record(abs(times[0] - start) <= time_tol, name + ": start time matches", {"actual": times[0], "expected": start})
            deltas = [b - a for a, b in zip(times, times[1:])]
            record(all(delta > 0 for delta in deltas), name + ": times are strictly increasing without duplicates")
            bad_steps = [i + 2 for i, delta in enumerate(deltas) if not (abs(delta - step) <= time_tol or
                         (allowed_event and i == len(deltas) - 1 and 0 < delta < step))]
            record(not bad_steps, name + ": time grid has every declared step except an allowed final event", {"bad_data_row_numbers": bad_steps[:20], "bad_row_count": len(bad_steps)})
            record(times[-1] + time_tol >= minimum_end, name + ": final time reaches minimum_end", {"actual": times[-1], "minimum_end": minimum_end, "data_rows": len(rows)})
            first = cfg.get("first_row_values")
            if first is None:
                report["not_covered"].append(name + ": initial physical values were not declared for checking.")
            else:
                if not isinstance(first, dict) or not first:
                    raise ContractError("first_row_values must map one-based column strings to numeric expected values")
                for key, expected in first.items():
                    if not isinstance(key, str) or not re.fullmatch(r"[1-9][0-9]*", key) or int(key) > width:
                        raise ContractError("first_row_values contains an invalid column")
                    expected = number(expected, "first_row_values")
                    actual = number(rows[0][int(key) - 1], "first row cell")
                    record(abs(actual - expected) <= tolerance, name + ": first-row value column " + key, {"actual": actual, "expected": expected, "absolute_tolerance": tolerance})
            maximum = cfg.get("max_value_exclusive")
            if maximum is None:
                report["not_covered"].append(name + ": no terminal strict numeric threshold was declared.")
            else:
                maximum = number(maximum, "max_value_exclusive")
                record(precision == "unrounded", name + ": strict threshold uses declared unrounded input precision")
                if precision == "unrounded":
                    actual = max(number(rows[-1][c - 1], "terminal value") for c in columns)
                    record(actual < maximum, name + ": last-row maximum is strictly below exclusive threshold", {"maximum": actual, "exclusive_threshold": maximum, "checked_columns": columns,
                           "comparison_uses_tolerance": False, "scope": "Last-row selected columns only; no whole-domain event certification"})
            samples = cfg.get("body_samples")
            if samples is None:
                report["not_covered"].append(name + ": no body CSV samples were declared; units and spatial/sample column selection are not checked.")
                continue
            object_fields(samples, {"table", "table_sha256", "time_key", "column_mappings", "tolerance"}, "body_samples")
            sample_path = bind({"path": samples.get("table"), "sha256": samples.get("table_sha256")}, name + " body CSV", {".csv"})
            sample_tol = number(samples.get("tolerance", tolerance), "body_samples.tolerance")
            if sample_tol < 0:
                raise ContractError("body_samples.tolerance must be nonnegative")
            key = samples.get("time_key")
            mappings = samples.get("column_mappings")
            if not isinstance(key, str) or not key or not isinstance(mappings, list) or not mappings:
                raise ContractError("body_samples needs a time_key and nonempty column_mappings")
            configured = []
            target_columns = set()
            for mapping in mappings:
                object_fields(mapping, {"body_column", "workbook_column", "scale", "offset"}, "column_mapping")
                body_column, target_col = mapping.get("body_column"), column(mapping.get("workbook_column"), "workbook_column")
                if not isinstance(body_column, str) or not body_column or target_col > width or target_col in target_columns:
                    raise ContractError("Body mappings require distinct valid workbook columns and nonempty CSV column names")
                target_columns.add(target_col)
                configured.append((body_column, target_col, number(mapping.get("scale", 1), "scale"), number(mapping.get("offset", 0), "offset")))
            with sample_path.open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream)
                fields = reader.fieldnames
                if not fields or len(fields) != len(set(fields)) or key not in fields or any(c not in fields for c, _, _, _ in configured):
                    raise ContractError("Body CSV needs unique headers and every declared time/value column")
                sample_rows = list(reader)
            record(bool(sample_rows), name + ": body CSV contains declared samples")
            sample_times = set()
            ordered = sorted((value, index) for index, value in enumerate(times))
            sorted_times = [value for value, _ in ordered]
            for index, body in enumerate(sample_rows, 1):
                if None in body or any(value is None for value in body.values()):
                    raise ContractError("Body CSV row width differs from its header")
                try:
                    t = float(body[key])
                except (ValueError, TypeError) as exc:
                    raise ContractError("Body sample times must be finite numbers") from exc
                if not math.isfinite(t):
                    raise ContractError("Body sample times must be finite numbers")
                record(t not in sample_times, name + ": body sample time is unique row " + str(index))
                sample_times.add(t)
                left = bisect.bisect_left(sorted_times, t - time_tol)
                right = bisect.bisect_right(sorted_times, t + time_tol)
                candidates = [ordered[n][1] for n in range(left, right)]
                record(len(candidates) == 1, name + ": body sample time exists unambiguously row " + str(index), {"time": t})
                if len(candidates) != 1:
                    continue
                row = rows[candidates[0]]
                for body_column, target_col, scale, offset in configured:
                    try:
                        expected = float(body[body_column]) * scale + offset
                    except (ValueError, TypeError) as exc:
                        raise ContractError("Body sample values must be finite numbers") from exc
                    expected = number(expected, "mapped body value")
                    actual = number(row[target_col - 1], "mapped workbook value")
                    record(abs(actual - expected) <= sample_tol, name + ": body sample row " + str(index) + " column " + str(target_col),
                           {"actual": actual, "mapped_expected": expected, "absolute_tolerance": sample_tol})
        report["not_covered"].extend(name + ": no time/value/sample contract declared for this sheet." for name in result.sheetnames if name not in names)
        record(set(names).issubset(template.sheetnames), "Declared sheet contracts belong to original template")
    except Exception as exc:
        report["errors"].append(str(exc) if isinstance(exc, ContractError) else "Input could not be inspected: " + type(exc).__name__)
    finally:
        for book in books:
            book.close()
        for key, path, before in frozen:
            try:
                current = confined(root, key)
                after = sha256(current)
            except (ContractError, OSError):
                after = None
            report["input_hashes"].append({"path": key, "before_sha256": before, "after_sha256": after, "unchanged": before == after})
            record(before == after, "Input unchanged after checks: " + key)
    report["status"] = "mechanical pass" if not report["errors"] else "failed"
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--contract", type=Path, required=True)
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--report", help="Optional NEW relative report JSON inside project-root; parent must exist")
    args = parser.parse_args()
    report = check(args.contract, args.project_root)
    if args.report:
        try:
            root = root_path(args.project_root)
            target = confined(root, args.report, existing=False)
            if target.exists() or not target.parent.is_dir():
                raise ContractError("Report must be a NEW file under an existing project-root directory")
            with target.open("x", encoding="utf-8", newline="\n") as stream:
                stream.write(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
        except (ContractError, OSError):
            report["errors"].append("Requested report path is unsafe, existing or unavailable; no overwrite was performed")
            report["status"] = "failed"
    print(json.dumps(report, ensure_ascii=False))
    return 0 if report["status"] == "mechanical pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
